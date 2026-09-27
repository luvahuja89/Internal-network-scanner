"""
Enhanced Multi-Sheet Excel Report Generator
Generates rich, styled .xlsx workbooks with:
1. Executive Summary & embedded Pie Chart
2. Detailed Asset Inventory with auto-filters & formatting
3. Prioritized Security Exposures
4. Delta / Change Audit (if previous scan exists)
"""

import datetime
from pathlib import Path
from typing import List, Dict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import PieChart, Reference


# Color Palette
NAVY_HEADER = "1F4E79"
SLATE_HEADER = "2C3E50"
WHITE = "FFFFFF"
LIGHT_GRAY = "F8F9FA"
BORDER_GRAY = "D6D8DB"

RISK_COLORS = {
    "CRITICAL": {"fill": "F8D7DA", "font": "721C24", "hex_chart": "DC3545"},
    "HIGH":     {"fill": "FFE5D0", "font": "A04000", "hex_chart": "FD7E14"},
    "MEDIUM":   {"fill": "FFF3CD", "font": "856404", "hex_chart": "FFC107"},
    "LOW":      {"fill": "D4EDDA", "font": "155724", "hex_chart": "28A745"},
    "INFO":     {"fill": "D1ECF1", "font": "0C5460", "hex_chart": "17A2B8"},
}


def create_thin_border():
    thin = Side(border_style="thin", color=BORDER_GRAY)
    return Border(top=thin, left=thin, right=thin, bottom=thin)


def auto_fit_columns(ws, min_width=12, max_width=50):
    """Auto-adjusts column widths based on content length."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = str(cell.value or '')
            if '\n' in val:
                val = max(val.split('\n'), key=len)
            max_len = max(max_len, len(val))
        adjusted = max(min(max_len + 3, max_width), min_width)
        ws.column_dimensions[col_letter].width = adjusted


def generate_excel_report(hosts_data: List[Dict], delta: Dict, target_display: str, output_path: Path):
    """
    Creates an executive multi-sheet Excel workbook (.xlsx).
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 1. Sheet 1: Executive Summary
    create_executive_summary_sheet(wb, hosts_data, delta, target_display)

    # 2. Sheet 2: Detailed Asset Inventory
    create_detailed_inventory_sheet(wb, hosts_data)

    # 3. Sheet 3: Prioritized Security Exposures
    create_security_exposures_sheet(wb, hosts_data)

    # 4. Sheet 4: Delta Audit
    if delta and (delta.get("new_hosts") or delta.get("removed_hosts") or delta.get("port_changes")):
        create_delta_audit_sheet(wb, delta)

    wb.save(output_path)


def create_executive_summary_sheet(wb, hosts_data: List[Dict], delta: Dict, target_display: str):
    ws = wb.create_sheet(title="Executive Summary")
    ws.views.sheetView[0].showGridLines = True
    thin_border = create_thin_border()

    # Title Banner
    ws.merge_cells("A1:G1")
    ws["A1"] = "🛡️ INTERNAL NETWORK & ASSET DISCOVERY — EXECUTIVE SUMMARY"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # Scan Metadata Box
    metadata = [
        ("Target Subnets / VLANs:", target_display),
        ("Scan Execution Time:", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total Live Assets Discovered:", len(hosts_data)),
        ("Total Accessible Services / Ports:", sum(len(h.get("ports", [])) for h in hosts_data)),
    ]

    for idx, (k, v) in enumerate(metadata, start=3):
        ws[f"A{idx}"] = k
        ws[f"A{idx}"].font = Font(name="Calibri", size=11, bold=True, color=SLATE_HEADER)
        ws[f"B{idx}"] = str(v)
        ws[f"B{idx}"].font = Font(name="Calibri", size=11, bold=False)
        ws[f"A{idx}"].border = thin_border
        ws[f"B{idx}"].border = thin_border

    # Calculate Risk Breakdown
    risk_counts = {"CRITICAL": 0, "HIGH": 0, "MEDIUM": 0, "LOW": 0, "INFO": 0}
    for h in hosts_data:
        r = h.get("overall_risk", "INFO")
        risk_counts[r] = risk_counts.get(r, 0) + 1

    # Calculate Category Breakdown
    cat_counts = {}
    for h in hosts_data:
        cat = h.get("category_label") or h.get("category", "General Host")
        cat_counts[cat] = cat_counts.get(cat, 0) + 1

    # Table 1: Risk Level Distribution (Rows 8 - 14)
    ws["A8"] = "Risk Severity"
    ws["B8"] = "Host Count"
    ws["A8"].font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    ws["B8"].font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    ws["A8"].fill = PatternFill(start_color=SLATE_HEADER, end_color=SLATE_HEADER, fill_type="solid")
    ws["B8"].fill = PatternFill(start_color=SLATE_HEADER, end_color=SLATE_HEADER, fill_type="solid")
    ws["A8"].alignment = Alignment(horizontal="left")
    ws["B8"].alignment = Alignment(horizontal="center")

    risk_row_start = 9
    for i, (risk_label, count) in enumerate(risk_counts.items()):
        row = risk_row_start + i
        ws[f"A{row}"] = risk_label
        ws[f"B{row}"] = count
        
        style = RISK_COLORS.get(risk_label, {})
        if style:
            ws[f"A{row}"].fill = PatternFill(start_color=style["fill"], end_color=style["fill"], fill_type="solid")
            ws[f"A{row}"].font = Font(name="Calibri", size=11, bold=True, color=style["font"])
        ws[f"B{row}"].font = Font(name="Calibri", size=11, bold=True)
        ws[f"B{row}"].alignment = Alignment(horizontal="center")
        ws[f"A{row}"].border = thin_border
        ws[f"B{row}"].border = thin_border

    # Pie Chart for Risk Distribution
    pie = PieChart()
    pie.title = "Asset Risk Distribution"
    labels = Reference(ws, min_col=1, min_row=risk_row_start, max_row=risk_row_start + len(risk_counts) - 1)
    data = Reference(ws, min_col=2, min_row=8, max_row=risk_row_start + len(risk_counts) - 1)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.width = 14
    pie.height = 7.5
    ws.add_chart(pie, "D3")

    # Table 2: Asset Categories Breakdown (Rows 16+)
    cat_row_start = 16
    ws[f"A{cat_row_start}"] = "Asset Classification / Category"
    ws[f"B{cat_row_start}"] = "Count"
    ws[f"A{cat_row_start}"].font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    ws[f"B{cat_row_start}"].font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    ws[f"A{cat_row_start}"].fill = PatternFill(start_color=SLATE_HEADER, end_color=SLATE_HEADER, fill_type="solid")
    ws[f"B{cat_row_start}"].fill = PatternFill(start_color=SLATE_HEADER, end_color=SLATE_HEADER, fill_type="solid")

    for i, (cat_name, count) in enumerate(sorted(cat_counts.items(), key=lambda x: x[1], reverse=True)):
        row = cat_row_start + 1 + i
        ws[f"A{row}"] = cat_name
        ws[f"B{row}"] = count
        ws[f"A{row}"].border = thin_border
        ws[f"B{row}"].border = thin_border
        ws[f"B{row}"].alignment = Alignment(horizontal="center")
        if i % 2 == 1:
            ws[f"A{row}"].fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
            ws[f"B{row}"].fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")

    auto_fit_columns(ws, min_width=18, max_width=45)


def create_detailed_inventory_sheet(wb, hosts_data: List[Dict]):
    ws = wb.create_sheet(title="Detailed Asset Inventory")
    ws.views.sheetView[0].showGridLines = True
    thin_border = create_thin_border()

    headers = [
        "IP Address", "Hostname", "Overall Risk", "Asset Classification",
        "MAC Address", "MAC Vendor", "Port", "Proto", "Service", "Version / Banner",
        "HTTP Title", "HTTP Server", "TLS Certificate CN", "TLS Expiry", "Port Risk", "Port Risk Findings"
    ]

    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    row_idx = 2
    for h in hosts_data:
        ip = h.get("ip", "")
        hostname = h.get("hostname", "N/A")
        overall_risk = h.get("overall_risk", "INFO")
        category = h.get("category_label") or h.get("category", "General Host")
        mac = h.get("mac_address", "N/A")
        vendor = h.get("mac_vendor", "")
        ports = h.get("ports", [])

        if not ports:
            row_data = [
                ip, hostname, overall_risk, category,
                mac, vendor, "N/A", "N/A", "N/A", "",
                "", "", "", "", "INFO", "Host active; no open ports in wellknown_os profile"
            ]
            ws.append(row_data)
            style_inventory_row(ws, row_idx, overall_risk, "INFO", thin_border, is_even=(row_idx % 2 == 0))
            row_idx += 1
        else:
            for p in ports:
                tls = p.get("tls_cert", {})
                tls_cn = tls.get("subject_cn", "")
                tls_exp = tls.get("expiry_date", "")
                p_risk = p.get("risk", "INFO")
                p_findings = p.get("risk_desc", "")

                row_data = [
                    ip, hostname, overall_risk, category,
                    mac, vendor, p.get("port"), p.get("protocol", "tcp"), p.get("service", ""), p.get("version", ""),
                    p.get("http_title", ""), p.get("http_server", ""), tls_cn, tls_exp, p_risk, p_findings
                ]
                ws.append(row_data)
                style_inventory_row(ws, row_idx, overall_risk, p_risk, thin_border, is_even=(row_idx % 2 == 0))
                row_idx += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    auto_fit_columns(ws, min_width=10, max_width=45)


def style_inventory_row(ws, row_idx, overall_risk, port_risk, border, is_even=False):
    for col_idx in range(1, 17):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.border = border
        cell.font = Font(name="Calibri", size=10)
        
        # Center align short columns
        if col_idx in [1, 7, 8, 14]:
            cell.alignment = Alignment(horizontal="center")
        elif col_idx in [3, 15]: # Risk Columns
            cell.alignment = Alignment(horizontal="center")
            r_val = cell.value
            if r_val in RISK_COLORS:
                cell.fill = PatternFill(start_color=RISK_COLORS[r_val]["fill"], end_color=RISK_COLORS[r_val]["fill"], fill_type="solid")
                cell.font = Font(name="Calibri", size=10, bold=True, color=RISK_COLORS[r_val]["font"])
        elif is_even:
            cell.fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")


def create_security_exposures_sheet(wb, hosts_data: List[Dict]):
    ws = wb.create_sheet(title="Prioritized Exposures")
    ws.views.sheetView[0].showGridLines = True
    thin_border = create_thin_border()

    headers = [
        "Severity", "Target IP", "Hostname", "Asset Classification",
        "Port / Service", "Technical Exposure & Risk Finding", "Recommended Remediation Action"
    ]
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=SLATE_HEADER, end_color=SLATE_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    remediation_guide = {
        2375: "Enforce TLS mutual authentication on port 2376 or bind Docker daemon strictly to unix:///var/run/docker.sock.",
        6379: "Require strong password in redis.conf ('requirepass'), bind to localhost/private VLAN, or enable ACLs.",
        9200: "Enable OpenSearch/Elastic Security plugin with TLS encryption, enforce RBAC, and disable public cluster binding.",
        10250: "Enable Kubelet authentication (--anonymous-auth=false) and enforce webhook authorization.",
        445:  "Restrict SMB access with network ACLs / host firewalls; disable SMBv1 and enforce SMB signing.",
        135:  "Filter Microsoft RPC (port 135) from untrusted subnets; restrict to domain controllers and management jump hosts.",
        3389: "Restrict RDP access behind VPN / Guacamole jump box; enable Network Level Authentication (NLA) and MFA.",
        23:   "Disable Telnet immediately; migrate administrative access to SSH with key-based authentication."
    }

    row_idx = 2
    # Collect all findings with risk CRITICAL, HIGH, MEDIUM
    exposure_list = []
    for h in hosts_data:
        ip = h.get("ip", "")
        hostname = h.get("hostname", "N/A")
        cat = h.get("category_label") or h.get("category", "General Host")
        for p in h.get("ports", []):
            risk = p.get("risk", "INFO")
            if risk in ["CRITICAL", "HIGH", "MEDIUM"]:
                port_num = p.get("port")
                remediation = remediation_guide.get(port_num, "Review service access control lists, enforce authentication, and restrict exposure to required VLANs.")
                exposure_list.append({
                    "risk": risk,
                    "ip": ip,
                    "hostname": hostname,
                    "category": cat,
                    "port_svc": f"{port_num}/{p.get('protocol','tcp')} ({p.get('service','')})",
                    "finding": p.get("risk_desc", ""),
                    "remediation": remediation
                })

    # Sort exposures: CRITICAL first, then HIGH, then MEDIUM
    risk_order = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2}
    exposure_list.sort(key=lambda x: risk_order.get(x["risk"], 3))

    for exp in exposure_list:
        row_data = [
            exp["risk"], exp["ip"], exp["hostname"], exp["category"],
            exp["port_svc"], exp["finding"], exp["remediation"]
        ]
        ws.append(row_data)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            if col_idx == 1:
                r_val = exp["risk"]
                cell.alignment = Alignment(horizontal="center")
                if r_val in RISK_COLORS:
                    cell.fill = PatternFill(start_color=RISK_COLORS[r_val]["fill"], end_color=RISK_COLORS[r_val]["fill"], fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color=RISK_COLORS[r_val]["font"])
            elif col_idx in [2, 5]:
                cell.alignment = Alignment(horizontal="center")
        row_idx += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    auto_fit_columns(ws, min_width=12, max_width=55)


def create_delta_audit_sheet(wb, delta: Dict):
    ws = wb.create_sheet(title="Delta Changes Audit")
    ws.views.sheetView[0].showGridLines = True
    thin_border = create_thin_border()

    headers = ["Change Event", "Target IP", "Hostname", "Event Details & Observations"]
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    row_idx = 2
    for h in delta.get("new_hosts", []):
        ws.append(["🚨 NEW ASSET DISCOVERED", h.get("ip"), h.get("hostname", "N/A"), f"New live asset detected on network (Category: {h.get('category','UNKNOWN')})"])
        ws.cell(row=row_idx, column=1).fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
        ws.cell(row=row_idx, column=1).font = Font(name="Calibri", size=10, bold=True, color="721C24")
        for c in range(1, 5):
            ws.cell(row=row_idx, column=c).border = thin_border
        row_idx += 1

    for h in delta.get("removed_hosts", []):
        ws.append(["⚪ DECOMMISSIONED / OFFLINE", h.get("ip"), h.get("hostname", "N/A"), "Host previously active but did not respond in current scan"])
        for c in range(1, 5):
            ws.cell(row=row_idx, column=c).border = thin_border
        row_idx += 1

    for pc in delta.get("port_changes", []):
        ws.append(["⚠️ PORT STATE CHANGE", pc.get("ip"), pc.get("hostname", "N/A"), f"Port {pc.get('port')}/{pc.get('protocol')} transitioned to {pc.get('action')}"])
        for c in range(1, 5):
            ws.cell(row=row_idx, column=c).border = thin_border
        row_idx += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    auto_fit_columns(ws, min_width=14, max_width=60)
